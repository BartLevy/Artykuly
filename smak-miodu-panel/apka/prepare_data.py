"""Wczytanie danych ślepego panelu (badania-2.xlsx, arkusze „Badania” i „Wyniki”).

Braki odpowiedzi pozostają brakami (None) — analiza na dostępnych danych.
"""

from dataclasses import asdict
from numbers import Number
from pathlib import Path

from openpyxl import Workbook, load_workbook

from data_model import DataModel

def get_excel_sheet(path: str, sheet_name: str | None = None):
    wb = load_workbook(path, data_only=True)
    return wb[sheet_name] if sheet_name else wb.active


def parse_number_from_cell(cell) -> int | None:
    """'2 - 25-45' -> 2, '3- za silna' -> 3, puste -> None."""
    if cell is None:
        return None
    s = str(cell).strip()
    if not s:
        return None
    try:
        return int(s[0])
    except ValueError:
        return None


def load_and_merge(fname: str) -> list[DataModel]:
    badania = get_excel_sheet(fname, "Badania")
    wyniki = get_excel_sheet(fname, "Wyniki")

    badania_by_id: dict[int, tuple] = {
        int(row[0]): row for row in badania.iter_rows(values_only=True)
        if isinstance(row[0], Number)
    }

    result = []
    for row_res in wyniki.iter_rows(values_only=True):
        if not isinstance(row_res[0], Number) or not row_res[1]:
            continue
        row = badania_by_id.get(int(row_res[0]))
        if row is None:
            continue
        result.append(DataModel(
            id=int(row[0]),
            date=row[1],
            age=parse_number_from_cell(row[2]),
            how_often=parse_number_from_cell(row[3]),
            gender=row[4] if row[4] in ("k", "m") else None,
            lot=row_res[1],
            overall_rate=row_res[2],
            sweetness=parse_number_from_cell(row_res[3]),
            acidity=parse_number_from_cell(row_res[4]),
            intensity=parse_number_from_cell(row_res[5]),
            is_typical={"t": 1, "n": 0}.get(row_res[6]),
            comment=row_res[7],
        ))
    return result


def save_to_excel(records: list[DataModel], output_path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(list(asdict(DataModel()).keys()))
    for dm in records:
        ws.append(list(asdict(dm).values()))
    wb.save(output_path)
