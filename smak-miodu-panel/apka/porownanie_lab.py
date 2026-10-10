"""Zestawienie ślepego panelu konsumenckiego z niezależną oceną organoleptyczną
czterech laboratoriów (zlecenie ODER 23/2026, sierpień 2026).

Opisy zapachu i smaku są czytane wprost z arkusza porównawczego i kodowane
regułowo na trzystopniową skalę typowości:
    2 = „typowy dla miodu / pożytku”,
    1 = „lekko / słabo typowy”,
    0 = „nietypowy dla miodu”.
Wynik laboratorium dla próbki = minimum z kodów zapachu i smaku (o ile podano
kwalifikator typowości). Zastrzeżenie organoleptyczne = wynik 0.
Laboratorium w Puławach nie stosuje kwalifikatora typowości (metoda
PN-88/A-77626, wyłącznie deskryptory), dlatego jego opisy raportowane są
osobno i nie wchodzą do miar zgodności.
"""

import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy import stats

XLSX = Path(__file__).resolve().parent.parent / "wyniki" / "Analiza_porownawcza_ODER_23_2026.xlsx"

# kod partii w arkuszu laboratoryjnym -> symbol próbki
LAB_CODE_TO_P = {
    "250906": "P1", "25C062101": "P2", "4/02/2029": "P3", "CQ23120101-2": "P4",
    "CQ25111501-1": "P5", "02/09": "P6", "21/05": "P7",
}
# kod partii w danych panelu -> symbol próbki
PANEL_LOT_TO_P = {
    "1-250906": "P1", "2-25c062101": "P2", "3-4/02/2020": "P3",
    "4-cq23120101-2": "P4", "5-cq25111501-1": "P5", "6-02/09": "P6",
    "7-4/21/05": "P7",
}
SAMPLES = ["P1", "P2", "P3", "P4", "P5", "P6", "P7"]
RATED_LABS = ["QSI", "Eurofins", "Intertek"]
LAB_NO = {"QSI": 1, "Eurofins": 2, "Intertek": 3, "Puławy": 4}

# kolory wg liczby laboratoriów z zastrzeżeniem (0, 1, 2, 3)
FLAG_CLR = {0: "#2874A6", 1: "#7F8C8D", 2: "#CA6F1E", 3: "#CA6F1E"}
FLAG_LABEL = {0: "bez zastrzeżeń laboratoriów (0/3)",
              1: "zastrzeżenie jednego laboratorium (1/3)",
              3: "zastrzeżenie wszystkich laboratoriów (3/3)"}


# --------------------------------------------------------------------------- #
#  Wczytanie i kodowanie opisów laboratoryjnych
# --------------------------------------------------------------------------- #

def typicality_code(text: str | None) -> int | None:
    if not text:
        return None
    t = str(text).lower()
    if "nietypow" in t:
        return 0
    if re.search(r"(lekko|słabo)\s+typow", t):
        return 1
    if "typow" in t:
        return 2
    return None


def load_lab_descriptions(path: Path = XLSX) -> pd.DataFrame:
    ws = load_workbook(path, data_only=True)["Macierz pełna"]
    rows = list(ws.iter_rows(values_only=True))
    codes_row = next(r for r in rows if r[0] == "Parametr")
    labs_row = rows[rows.index(codes_row) + 1]
    smell = next(r for r in rows if r[0] == "Zapach")
    taste = next(r for r in rows if r[0] == "Smak")

    out, code = [], None
    for j in range(1, len(codes_row)):
        if codes_row[j] is not None:
            code = str(codes_row[j])
        lab = labs_row[j]
        if lab is None or code not in LAB_CODE_TO_P:
            continue
        out.append({"P": LAB_CODE_TO_P[code], "lab": lab,
                    "smell": smell[j], "taste": taste[j]})
    d = pd.DataFrame(out)
    d["smell_code"] = d["smell"].map(typicality_code)
    d["taste_code"] = d["taste"].map(typicality_code)
    d["score"] = d[["smell_code", "taste_code"]].min(axis=1, skipna=True)
    d["flag"] = (d["score"] == 0).astype(int)
    return d


# --------------------------------------------------------------------------- #
#  Miary
# --------------------------------------------------------------------------- #

def wilson_ci(k: int, n: int, z: float = 1.96) -> tuple[float, float]:
    p = k / n
    den = 1 + z ** 2 / n
    c = (p + z ** 2 / (2 * n)) / den
    h = z * np.sqrt(p * (1 - p) / n + z ** 2 / (4 * n ** 2)) / den
    return c - h, c + h


def fleiss_kappa(table: np.ndarray) -> float:
    """table: n_obiektów × n_kategorii (liczba oceniających w kategorii)."""
    n_raters = table.sum(axis=1)[0]
    p_j = table.sum(axis=0) / table.sum()
    P_i = ((table ** 2).sum(axis=1) - n_raters) / (n_raters * (n_raters - 1))
    P_bar, P_e = P_i.mean(), (p_j ** 2).sum()
    return float((P_bar - P_e) / (1 - P_e)) if P_e < 1 else float("nan")


def cochran_q(x: np.ndarray) -> tuple[float, float, int]:
    """x: n_respondentów × k_próbek (0/1). Zwraca Q, p, df."""
    k = x.shape[1]
    col, row, total = x.sum(axis=0), x.sum(axis=1), x.sum()
    q = (k - 1) * (k * (col ** 2).sum() - total ** 2) / (k * total - (row ** 2).sum())
    return float(q), float(stats.chi2.sf(q, k - 1)), k - 1


def paired_contrast(wide: pd.DataFrame, a: list[str], b: list[str]) -> dict:
    """Różnica średnich (a − b) per respondent: średnia, 95% CI, Wilcoxon, d."""
    diff = wide[a].mean(axis=1) - wide[b].mean(axis=1)
    sd = diff.std(ddof=1)
    half = 1.96 * sd / np.sqrt(len(diff))
    return {"mean": diff.mean(), "lo": diff.mean() - half, "hi": diff.mean() + half,
            "p": stats.wilcoxon(diff[diff != 0]).pvalue, "d": diff.mean() / sd}


def _f(x: float, nd: int = 3) -> str:
    return f"{x:.{nd}f}".replace(".", "{,}")


def _p(p: float) -> str:
    return "< 0{,}001" if p < 0.001 else "= " + _f(p, 3)


def _pt(p: float) -> str:
    return "$<$0{,}001" if p < 0.001 else _f(p, 3)


def _pct(x: float) -> str:
    return f"{100 * x:.1f}".replace(".", "{,}")


# --------------------------------------------------------------------------- #
#  Główne obliczenia
# --------------------------------------------------------------------------- #

def compute_lab_comparison(df: pd.DataFrame, out_dir: str = "..") -> dict:
    d = df.copy()
    d["P"] = d["lot"].map(PANEL_LOT_TO_P)
    labs = load_lab_descriptions()
    r: dict = {}

    # --- panel: profil próbek ---
    g = d.groupby("P")
    panel = pd.DataFrame({
        "n": g["is_typical"].count(),
        "k": g["is_typical"].sum().astype(int),
        "typ": g["is_typical"].mean(),
        "overall": g["overall_rate"].mean(),
        "sweet": g["sweetness"].mean(),
        "acid": g["acidity"].mean(),
        "intens": g["intensity"].mean(),
        "sweet3": g["sweetness"].apply(lambda s: (s.dropna() == 3).mean()),
    }).loc[SAMPLES]
    panel[["lo", "hi"]] = [wilson_ci(int(k), int(n)) for k, n in zip(panel.k, panel.n)]
    panel["p_binom"] = [stats.binomtest(int(k), int(n), 0.5).pvalue
                        for k, n in zip(panel.k, panel.n)]

    # --- laboratoria: wynik i zastrzeżenia ---
    rated = labs[labs.lab.isin(RATED_LABS)]
    score = rated.pivot(index="P", columns="lab", values="score").loc[SAMPLES, RATED_LABS]
    flag = rated.pivot(index="P", columns="lab", values="flag").loc[SAMPLES, RATED_LABS]
    lab_mean = score.mean(axis=1)
    n_flags = flag.sum(axis=1).astype(int)
    lab_majority = (n_flags >= 2).astype(int)          # 1 = laboratoria kwestionują
    panel_majority = (panel["typ"] < 0.5).astype(int)  # 1 = większość panelu: nietypowy

    # zgodność między laboratoriami (3 oceniające)
    cnt_bin = np.column_stack([(flag == 0).sum(axis=1), (flag == 1).sum(axis=1)])
    cnt_ord = np.column_stack([(score == c).sum(axis=1) for c in (0, 1, 2)])
    r["kappa_bin"] = _f(fleiss_kappa(cnt_bin), 2)
    r["kappa_ord"] = _f(fleiss_kappa(cnt_ord), 2)
    r["n_full_agree_bin"] = int(((n_flags == 0) | (n_flags == 3)).sum())
    r["n_full_agree_ord"] = int((score.nunique(axis=1) == 1).sum())

    # panel vs laboratoria (poziom próbek, n = 7)
    rho, p_rho = stats.spearmanr(panel["typ"], lab_mean)
    r["rho_panel_lab"], r["p_rho_panel_lab"] = _f(rho, 2), _p(p_rho)
    rho_o, p_rho_o = stats.spearmanr(panel["overall"], lab_mean)
    r["rho_overall_lab"], r["p_rho_overall_lab"] = _f(rho_o, 2), _p(p_rho_o)
    wo3 = [p for p in SAMPLES if p != "P3"]
    rho3, p_rho3 = stats.spearmanr(panel.loc[wo3, "typ"], lab_mean[wo3])
    r["rho_panel_lab_wo3"], r["p_rho_panel_lab_wo3"] = _f(rho3, 2), _p(p_rho3)
    r["agree_panel_lab"] = int((panel_majority == lab_majority).sum())

    # zdolność pojedynczej oceny panelu do odtworzenia werdyktu laboratoriów
    from analiza import auc
    dt = d.dropna(subset=["is_typical"])
    y_lab = dt["P"].map(lab_majority)
    atyp = 1 - dt["is_typical"].astype(float)
    r["auc_lab"] = _f(auc(atyp.to_numpy(), y_lab.to_numpy()))
    r["fp_lab_pct"] = _pct(atyp[y_lab == 0].mean())       # „nietypowy”, choć laby bez zastrzeżeń
    r["fn_lab_pct"] = _pct(1 - atyp[y_lab == 1].mean())   # „typowy”, choć laby kwestionują

    # Cochran Q: czy panel różnicuje próbki (wszystkie / bez P3)
    wide = d.pivot_table(index="id", columns="P", values="is_typical").dropna()
    q7, pq7, df7 = cochran_q(wide[SAMPLES].to_numpy())
    q6, pq6, df6 = cochran_q(wide[wo3].to_numpy())
    r["n_complete_typ"] = len(wide)
    r.update(q_all=_f(q7, 2), p_q_all=_p(pq7), df_q_all=df7,
             q_rest=_f(q6, 2), p_q_rest=_p(pq6), df_q_rest=df6)
    # Friedman dla ogólnej akceptacji (skala 1–5)
    wide_o = d.pivot_table(index="id", columns="P", values="overall_rate").dropna()
    r["n_complete_overall"] = len(wide_o)
    fr = stats.friedmanchisquare(*[wide_o[p] for p in SAMPLES])
    fr6 = stats.friedmanchisquare(*[wide_o[p] for p in wo3])
    r.update(fr_all=_f(fr.statistic, 2), p_fr_all=_p(fr.pvalue),
             fr_rest=_f(fr6.statistic, 2), p_fr_rest=_p(fr6.pvalue))

    # kontrast: próbki jednogłośnie zakwestionowane przez laboratoria vs bez zastrzeżeń
    flagged = [p for p in SAMPLES if n_flags[p] == 3]
    clean = [p for p in SAMPLES if n_flags[p] == 0]
    clean_wo = [p for p in clean if p != "P3"]
    c = paired_contrast(wide, clean, flagged)
    c2 = paired_contrast(wide, clean_wo, flagged)
    co = paired_contrast(wide_o, clean, flagged)
    r.update(
        lab_flagged=", ".join(flagged), lab_clean=", ".join(clean),
        lab_clean_wo=", ".join(clean_wo),
        typ_flagged=_f(d[d.P.isin(flagged)]["is_typical"].mean()),
        typ_clean=_f(d[d.P.isin(clean)]["is_typical"].mean()),
        diff_flag=_f(c["mean"]), diff_flag_lo=_f(c["lo"]), diff_flag_hi=_f(c["hi"]),
        p_diff_flag=_p(c["p"]), d_diff_flag=_f(c["d"], 2),
        diff_flag_wo=_f(c2["mean"]), p_diff_flag_wo=_p(c2["p"]),
        diff_overall_flag=_f(co["mean"], 2), p_diff_overall_flag=_p(co["p"]),
    )

    # wartości per próbka do tekstu
    for p in SAMPLES:
        r[f"typ_{p}"] = _f(panel.loc[p, "typ"])
        r[f"sweet3_{p}"] = _pct(panel.loc[p, "sweet3"])
        r[f"pbin_{p}"] = _p(panel.loc[p, "p_binom"])
    r["n_sig_binom"] = int((panel["p_binom"] < 0.05).sum())

    # --- tabela: profil sensoryczny w panelu ---
    r["profile_rows"] = "\n".join(
        f"{p} & {int(panel.loc[p, 'n'])} & "
        f"{_f(panel.loc[p, 'overall'], 2)} & {_f(panel.loc[p, 'typ'])} & "
        f"{_f(panel.loc[p, 'sweet'], 2)} & {_f(panel.loc[p, 'acid'], 2)} & "
        f"{_f(panel.loc[p, 'intens'], 2)} \\\\"
        for p in SAMPLES)

    # --- tabela: panel vs laboratoria ---
    def sym(v):
        return {0: "N", 1: "L", 2: "T"}.get(v, "--") if pd.notna(v) else "--"

    r["panel_lab_rows"] = "\n".join(
        f"{p} & {int(panel.loc[p, 'n'])} & {_f(panel.loc[p, 'typ'])} & "
        f"{_f(panel.loc[p, 'lo'], 2)}--{_f(panel.loc[p, 'hi'], 2)} & "
        f"{_pt(panel.loc[p, 'p_binom'])} & "
        + " & ".join(sym(score.loc[p, lab]) for lab in RATED_LABS)
        + f" & {n_flags[p]}/3 \\\\"
        for p in SAMPLES)

    save_panel_lab_figure(panel, n_flags, out_dir)
    save_contrast_figure(wide, clean, flagged, out_dir)
    return r


# --------------------------------------------------------------------------- #
#  Wykresy
# --------------------------------------------------------------------------- #

def _style() -> None:
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10,
                         "axes.spines.top": False, "axes.spines.right": False})


def _save(fig, out_dir: str, stem: str) -> None:
    p = Path(out_dir)
    fig.savefig(p / f"{stem}.pdf", bbox_inches="tight")
    fig.savefig(p / f"{stem}.png", bbox_inches="tight", dpi=150)
    plt.close(fig)
    print(f"  Zapisano {stem}.pdf / .png")


def save_panel_lab_figure(panel: pd.DataFrame, n_flags: pd.Series,
                          out_dir: str = "..") -> None:
    """Typowość i ogólna akceptacja w panelu, kolor wg werdyktu laboratoriów."""
    _style()
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(11, 4.4))
    x = np.arange(len(SAMPLES))
    for i, p in enumerate(SAMPLES):
        clr = FLAG_CLR[n_flags[p]]
        y, lo, hi = panel.loc[p, ["typ", "lo", "hi"]]
        ax1.errorbar(i, y, yerr=[[y - lo], [hi - y]], fmt="o", color=clr,
                     ms=8, capsize=4, lw=1.4)
        ax1.text(i, 0.05, f"{n_flags[p]}/3", ha="center", va="center",
                 fontsize=9.5, fontweight="bold", color=clr)
        ax2.bar(i, panel.loc[p, "overall"], color=clr, width=0.6)
        ax2.text(i, panel.loc[p, "overall"] + 0.06, f"{panel.loc[p, 'overall']:.2f}".replace(".", ","),
                 ha="center", va="bottom", fontsize=8.5)
    labels = SAMPLES
    ax1.axhline(0.5, color="grey", ls="--", lw=0.8, zorder=0)
    ax1.set(xticks=x, ylim=(0, 1.0), title="Typowość smaku (95% CI Wilsona)",
            ylabel="Odsetek ocen „smak typowy”")
    ax1.set_xticklabels(labels, fontsize=10)
    ax1.text(-0.45, 0.12, "zastrzeżenia laboratoriów", fontsize=7.5, color="#555555")
    ax2.axhline(3.0, color="grey", ls="--", lw=0.8, zorder=0)
    ax2.set(xticks=x, ylim=(1, 5), title="Ogólna akceptacja",
            ylabel="Średnia ocena (1–5)")
    ax2.set_xticklabels(labels, fontsize=10)
    fig.tight_layout(rect=(0, 0.08, 1, 1))
    present = sorted({0 if v == 0 else (1 if v == 1 else 3) for v in n_flags})
    fig.legend(handles=[Line2D([], [], marker="o", ls="", color=FLAG_CLR[k],
                               label=FLAG_LABEL[k]) for k in present],
               loc="lower center", ncol=len(present), frameon=False, fontsize=9)
    _save(fig, out_dir, "fig_panel_lab")


def save_contrast_figure(wide: pd.DataFrame, clean: list[str], flagged: list[str],
                         out_dir: str = "..") -> None:
    """Średnia typowość per respondent: próbki bez zastrzeżeń vs zakwestionowane."""
    _style()
    xs, ys = wide[flagged].mean(axis=1), wide[clean].mean(axis=1)
    pts = pd.DataFrame({"x": xs.round(3), "y": ys.round(3)}).value_counts().reset_index(name="n")
    fig, ax = plt.subplots(figsize=(5, 5))
    ax.scatter(pts.x, pts.y, s=pts.n * 28, color="#2874A6", alpha=0.6, edgecolor="white")
    ax.plot([0, 1], [0, 1], ls="--", color="grey", lw=0.8)
    ax.set(xlim=(-0.05, 1.05), ylim=(-0.05, 1.05), aspect="equal",
           xlabel=f"Typowość próbek zakwestionowanych ({', '.join(flagged)})",
           ylabel=f"Typowość próbek bez zastrzeżeń ({', '.join(clean)})")
    for n in (1, 5, 10):
        ax.scatter([], [], s=n * 28, color="#2874A6", alpha=0.6, label=f"n = {n}")
    ax.legend(frameon=False, fontsize=8, loc="lower right", title="respondenci",
              title_fontsize=8)
    fig.tight_layout()
    _save(fig, out_dir, "fig_contrast_lab")


if __name__ == "__main__":
    from prepare_data import load_and_merge
    from analiza import to_dataframe
    res = compute_lab_comparison(to_dataframe(load_and_merge("badania-2.xlsx")))
    for k, v in res.items():
        print(k, "=", v)
