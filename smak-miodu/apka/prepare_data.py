from dataclasses import asdict, replace
from numbers import Number
from pathlib import Path
from statistics import median

from openpyxl import load_workbook
from openpyxl import Workbook

from data_model import DataModel

_NUMERIC_FIELDS = ["age", "how_often", "overall_rate", "sweetness", "acidity", "intensity"]
_CATEGORICAL_FIELDS = ["gender"]


def get_excel_sheet(path: str, sheet_name: str | None = None):
    wb = load_workbook(path, data_only=True)
    sheet = wb[sheet_name] if sheet_name else wb.active
    return sheet


def parse_number_from_cell(cell) -> int | None:
    if cell is None:
        return None
    s = str(cell).strip()
    if not s:
        return None
    try:
        return int(s[0])
    except ValueError:
        return None


def _distance(a: DataModel, b: DataModel) -> float:
    total, count = 0.0, 0
    for field in _NUMERIC_FIELDS:
        va, vb = getattr(a, field), getattr(b, field)
        if va is not None and vb is not None:
            total += abs(va - vb)
            count += 1
    for field in _CATEGORICAL_FIELDS:
        va, vb = getattr(a, field), getattr(b, field)
        if va is not None and vb is not None:
            total += 0 if va == vb else 1
            count += 1
    return total / count if count else float("inf")


def interpolate_missing_by_knn(records: list[DataModel], k: int = 3) -> list[DataModel]:
    """Fill None numeric fields using the median of the k most similar rows."""
    result = [replace(r) for r in records]

    for i, rec in enumerate(result):
        missing = [f for f in _NUMERIC_FIELDS if getattr(rec, f) is None]
        if not missing:
            continue

        others = [r for j, r in enumerate(records) if j != i]
        neighbors = sorted(others, key=lambda r: _distance(rec, r))[:k]

        for field in missing:
            values = [getattr(n, field) for n in neighbors if getattr(n, field) is not None]
            if values:
                setattr(result[i], field, int(round(median(values))))

    return result

def get_authorized(lot):
    #                                        zła
    if (lot in ["1-250906","3-4/02/2020","5-cq25111501-1"]):
        return 1
    else:
        return 0
        
def get_flower(lot):
    if (lot == "1-250906"): return "robinia"
    if (lot == "2-25c062101"): return "robinia"
    if (lot == "3-4/02/2020"): return "polyfloral"
    if (lot == "4-cq23120101-2"): return "polyfloral"
    if (lot == "5-cq25111501-1"): return "tilia"
    if (lot == "6-02/09"): return "polyfloral"
    if (lot == "7-4/21/05"): return "polyfloral"
    return "other"

def load_and_merge(fname):
    b = get_excel_sheet(fname, "Badania")
    w = get_excel_sheet(fname, "Wyniki")

    badania_by_id: dict[int, tuple] = {}
    for row in b.iter_rows(values_only=True):
        if row[0] is not None and isinstance(row[0], Number):
            badania_by_id[int(row[0])] = row

    result = []
    for row_res in w.iter_rows(values_only=True):
        if row_res[0] is None or not isinstance(row_res[0], Number):
            continue
        row = badania_by_id.get(int(row_res[0]))
        if row is None:
            continue

        dm = DataModel()
        dm.id = int(row[0])
        # dm.date = row[1] if row[1] is not None else "2026-07-01"
        dm.age = parse_number_from_cell(row[2])
        dm.how_often = parse_number_from_cell(row[3])
        dm.gender = "m" if row[4] == "m" else "k"
        if (dm.gender is None): dm.gender = "k"
        dm.lot = row_res[1]
        dm.overall_rate = row_res[2]
        dm.sweetness = parse_number_from_cell(row_res[3])
        dm.acidity = parse_number_from_cell(row_res[4])
        dm.intensity = parse_number_from_cell(row_res[5])
        dm.is_typical = 1 if row_res[6] == "t" else 0
        dm.authorized = get_authorized(row_res[1])
        dm.type_of_honey = get_flower(row_res[1])
        result.append(dm)

    return interpolate_missing_by_knn(result)


def save_to_excel(records: list[DataModel], output_path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(list(asdict(DataModel()).keys()))
    for dm in records:
        ws.append(list(asdict(dm).values()))
    wb.save(output_path)
